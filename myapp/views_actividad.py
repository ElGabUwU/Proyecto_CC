"""
Vistas para la gestión de actividades comunitarios del Consejo Comunal.
"""
from django.shortcuts import render, redirect, get_object_or_404, reverse
from django.contrib import messages
from django.core.paginator import Paginator
from django.db.models import Q, Sum, Count
from django.http import JsonResponse, HttpResponse
from django.utils import timezone
from django.views.decorators.http import require_POST, require_GET
from django.db import transaction
import json
from datetime import datetime, timedelta
import csv
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ✅ CORRECCIÓN: Eliminado 'Person' de las importaciones
from .models import (
    Actividad, ActividadParticipante, Habitante
)
from .forms import (
    ActividadForm, AsignarParticipanteForm
)
from .decorators import admin_required


# ============================================
# Vistas para Actividades
# ============================================

def actividad(request):
    """
    Vista para listar y buscar actividades comunitarios.
    """
    query = request.GET.get('q', '')
    estatus = request.GET.get('estatus', '')
    categoria = request.GET.get('categoria', '')
    
    actividad_list = Actividad.objects.filter(is_deleted=False).select_related('creado_por')
    
    if query:
        actividad_list = actividad_list.filter(
            Q(nombre_actividad__icontains=query) |
            Q(descripcion__icontains=query)
        )
    
    if estatus:
        actividad_list = actividad_list.filter(estatus=estatus)
    
    if categoria:
        actividad_list = actividad_list.filter(categoria_enfoque=categoria)
    
    actividad_list = actividad_list.annotate(num_participantes=Count('participantes'))
    actividad_list = actividad_list.order_by('-fecha_creacion')
    
    paginator = Paginator(actividad_list, 15)
    page_number = request.GET.get('page')
    actividad_page = paginator.get_page(page_number)

    conteos = actividad_list.aggregate(
        activo=Count('pk', filter=Q(estatus='activo')),
        cerrado=Count('pk', filter=Q(estatus='cerrado')),
        archivado=Count('pk', filter=Q(estatus='archivado')),
    )
    
    context = {
        'actividad': actividad_page,
        'query': query,
        'estatus': estatus,
        'categoria': categoria,
        'actividad_form': ActividadForm(user=request.user),
        'estatus_choices': Actividad.ESTATUS_CHOICES,
        'categoria_choices': Actividad.CATEGORIA_ENFOQUE_CHOICES,
        'count_activo': conteos['activo'],
        'count_cerrado': conteos['cerrado'],
        'count_archivado': conteos['archivado'],
    }
    
    return render(request, 'actividad/actividad_list.html', context)


def crear_actividad(request):
    if request.method == 'POST':
        form = ActividadForm(request.POST, user=request.user)
        if form.is_valid():
            actividad = form.save()
            messages.success(request, f'Actividad "{actividad.nombre_actividad}" creado exitosamente.')
            # Si es una petición AJAX, devolver respuesta JSON
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({
                    'success': True,
                    'message': f'Actividad "{actividad.nombre_actividad}" creado exitosamente.',
                    'redirect_url': str(reverse('detalle_actividad', kwargs={'pk': actividad.pk}))
                })
            return redirect('actividad')
        else:
            # Si es una petición AJAX, devolver los errores del formulario en JSON
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                errors = {}
                for field, error_list in form.errors.items():
                    field_label = form.fields[field].label if field in form.fields else field
                    errors[field] = [str(e) for e in error_list]
                return JsonResponse({
                    'success': False,
                    'errors': errors,
                    'message': 'Por favor corrija los errores en el formulario.'
                }, status=400)
            messages.error(request, 'Por favor corrija los errores en el formulario.')
    else:
        form = ActividadForm(user=request.user)
    
    context = {'form': form}
    return render(request, 'actividad/actividad_list.html', context)


def detalle_actividad(request, pk):
    """
    Vista para ver el detalle completo de un actividad.
    Incluye el listado de participantes registrados.
    """
    # ✅ CORRECCIÓN: Rutas adaptadas al modelo unificado
    actividad = get_object_or_404(
        Actividad.objects.filter(is_deleted=False).prefetch_related(
            'participantes_actividad__habitante',
            'participantes_actividad__habitante__familia'
        ).select_related('creado_por'),
        pk=pk
    )
    
    participantes = actividad.participantes_actividad.all()
    stats_genero = {}
    stats_familias = {}
    
    for p in participantes:
        # ✅ CORRECCIÓN: Acceso directo a campos de Habitante
        genero_display = p.habitante.get_genero_display()
        stats_genero[genero_display] = stats_genero.get(genero_display, 0) + 1
        
        jefe = p.habitante.familia.jefe_familia
        familia_nombre = f"{jefe.nombre} {jefe.apellido}" if jefe else "Sin definir"
        stats_familias[familia_nombre] = stats_familias.get(familia_nombre, 0) + 1
    
    context = {
        'actividad': actividad,
        'participantes': participantes,
        'asignar_form': AsignarParticipanteForm(actividad=actividad),
        'stats_genero': stats_genero,
        'stats_familias': stats_familias,
    }
    
    return render(request, 'actividad/actividad_detail.html', context)


def editar_actividad(request, pk):
    actividad = get_object_or_404(Actividad, pk=pk, is_deleted=False)

    if request.method == "POST":
        form = ActividadForm(request.POST, instance=actividad, user=request.user)
        if form.is_valid():
            actividad = form.save()
            messages.success(request, f'Actividad "{actividad.nombre_actividad}" actualizado exitosamente.')
            # Si es una petición AJAX, devolver respuesta JSON
            if request.headers.get("X-Requested-With") == "XMLHttpRequest":
                return JsonResponse({
                    "success": True,
                    "message": f'Actividad "{actividad.nombre_actividad}" actualizado exitosamente.',
                    "redirect_url": str(reverse("detalle_actividad", kwargs={"pk": actividad.pk}))
                })
            return redirect("detalle_actividad", pk=actividad.pk)
        else:
            # Si es una petición AJAX, devolver los errores del formulario en JSON
            if request.headers.get("X-Requested-With") == "XMLHttpRequest":
                errors = {}
                for field, error_list in form.errors.items():
                    field_label = form.fields[field].label if field in form.fields else field
                    errors[field] = [str(e) for e in error_list]
                return JsonResponse({
                    "success": False,
                    "errors": errors,
                    "message": "Por favor corrija los errores en el formulario."
                }, status=400)
            messages.error(request, "Por favor corrija los errores en el formulario.")
    else:
        form = ActividadForm(instance=actividad, user=request.user)

    context = {"form": form, "actividad": actividad}
    return render(request, "actividad/actividad_form.html", context)


@require_POST
def eliminar_actividad(request, pk):
    actividad = get_object_or_404(Actividad, pk=pk, is_deleted=False)
    nombre = actividad.nombre_actividad
    actividad.delete()
    messages.success(request, f'Actividad "{nombre}" eliminado exitosamente.')
    return redirect('actividad')


@require_GET
def api_actividad(request, pk):
    actividad = get_object_or_404(Actividad, pk=pk, is_deleted=False)
    
    data = {
        'id': actividad.id,
        'nombre_actividad': actividad.nombre_actividad,
        'fecha_inicio': actividad.fecha_inicio.strftime('%Y-%m-%d'),
        'fecha_fin': actividad.fecha_fin.strftime('%Y-%m-%d') if actividad.fecha_fin else '',
        'descripcion': actividad.descripcion,
        'categoria_enfoque': actividad.categoria_enfoque,
        'categoria_enfoque_display': actividad.get_categoria_enfoque_display(),
        'estatus': actividad.estatus,
        'estatus_display': actividad.get_estatus_display(),
        'fecha_creacion': actividad.fecha_creacion.strftime('%d/%m/%Y %H:%M'),
        'creado_por': actividad.creado_por.username if actividad.creado_por else '',
        'participantes_count': actividad.cantidad_participantes(),
    }
    
    return JsonResponse(data)


# ============================================
# Vistas para Gestión de Participantes
# ============================================

def buscar_habitantes_actividad(request):
    """
    API para buscar habitantes para registrar en una actividad.
    """
    query = request.GET.get('q', '')
    actividad_id = request.GET.get('actividad_id', '')
    
    if len(query) < 2:
        return JsonResponse({'habitantes': []})
    
    # ✅ CORRECCIÓN: Filtros y selects adaptados a Habitante
    habitantes = Habitante.objects.filter(
        is_deleted=False
    ).filter(
        Q(nombre__icontains=query) |
        Q(apellido__icontains=query) |
        Q(cedula__icontains=query)
    ).select_related('familia').order_by('nombre')
    
    if actividad_id:
        registrados = ActividadParticipante.objects.filter(
            actividad_id=actividad_id
        ).values_list('habitante_id', flat=True)
        habitantes = habitantes.exclude(id__in=registrados)

    habitantes = habitantes[:20]
    
    resultados = []
    for h in habitantes:
        jefe = h.familia.jefe_familia
        resultados.append({
            'id': h.id,
            'nombre': f"{h.nombre} {h.apellido}",
            'documento': h.cedula,
            'telefono': getattr(h, 'telefono', ''),
            'familia': f"{jefe.nombre} {jefe.apellido}" if jefe else "Sin definir",
            'parentesco': "Jefe de Familia" if h.es_jefe_familia else "Miembro",
            'genero': h.get_genero_display(),
        })
    
    return JsonResponse({'habitantes': resultados})

def asignar_participante(request, pk):
    actividad = get_object_or_404(Actividad, pk=pk, is_deleted=False)
    
    if request.method == 'POST':
        form = AsignarParticipanteForm(request.POST, actividad=actividad)
        if form.is_valid():
            participante = form.save()
            # ✅ CORRECCIÓN: Nombres directos del habitante
            nombre_completo = f"{participante.habitante.nombre} {participante.habitante.apellido}"
            messages.success(request, f'{nombre_completo} ha sido registrado/a exitosamente en la actividad.')
            return redirect('detalle_actividad', pk=actividad.pk)
        else:
            messages.error(request, 'Por favor corrija los errores en el formulario.')
    else:
        form = AsignarParticipanteForm(actividad=actividad)
    
    return render(request, 'actividad/actividad_detail.html', {'form': form, 'actividad': actividad})


@require_POST
def remover_participante(request, pk, habitante_id):
    actividad = get_object_or_404(Actividad, pk=pk, is_deleted=False)
    
    try:
        participante = ActividadParticipante.objects.get(
            actividad=actividad,
            habitante_id=habitante_id
        )
        nombre_completo = f"{participante.habitante.nombre} {participante.habitante.apellido}"
        participante.delete()
        messages.success(request, f'{nombre_completo} ha sido removido/a de la actividad exitosamente.')
    except ActividadParticipante.DoesNotExist:
        messages.error(request, 'El habitante no está registrado en esta actividad.')
    
    return redirect('detalle_actividad', pk=actividad.pk)


@require_GET
def api_participante(request, participante_id):
    participante = get_object_or_404(ActividadParticipante, id=participante_id)
    
    data = {
        'id': participante.id,
        'actividad': participante.actividad.id,
        'actividad_nombre': participante.actividad.nombre_actividad,
        'habitante': participante.habitante.id,
        'habitante_nombre': participante.nombre_habitante,
        'habitante_documento': participante.documento_habitante,
        'fecha_registro': participante.fecha_registro.strftime('%d/%m/%Y %H:%M'),
        'observaciones': participante.observaciones or '',
    }
    
    return JsonResponse(data)


# ============================================
# Vistas para Exportación
# ============================================

def exportar_participantes_actividad(request, pk):
    """
    Vista para exportar la lista de participantes de una actividad a CSV.
    """
    actividad = get_object_or_404(Actividad, pk=pk, is_deleted=False)
    
    # ✅ CORRECCIÓN: Rutas de relación corregidas
    participantes = ActividadParticipante.objects.filter(
        actividad=actividad
    ).select_related(
        'habitante',
        'habitante__familia'
    ).order_by('habitante__nombre')
    
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="actividad_{actividad.id}_{datetime.now().strftime("%Y%m%d")}.csv"'
    
    writer = csv.writer(response, delimiter=';', quoting=csv.QUOTE_NONNUMERIC)
    writer.writerow([
        'N°', 'Cédula', 'Nombre Completo', 'Teléfono', 'Email',
        'Género', 'Familia', 'Parentesco', 'Fecha Registro', 'Observaciones'
    ])
    
    for i, p in enumerate(participantes, 1):
        h = p.habitante
        jefe = h.familia.jefe_familia
        
        writer.writerow([
            i,
            h.cedula,
            f"{h.nombre} {h.apellido}",
            getattr(h, 'telefono', ''),
            getattr(h, 'email', ''),
            h.get_genero_display(),
            f"{jefe.nombre} {jefe.apellido}" if jefe else "",
            "Jefe de Familia" if h.es_jefe_familia else "Miembro",
            p.fecha_registro.strftime('%d/%m/%Y %H:%M'),
            p.observaciones or '',
        ])
    
    return response


def exportar_participantes_actividad_excel(request, pk):
    """
    Genera un archivo Excel profesional (.xlsx) en memoria RAM con los 
    participantes de una actividad comunitaria, usando el mismo estilizado 
    que los reportes demográficos.
    """
    actividad = get_object_or_404(Actividad, pk=pk, is_deleted=False)
    
    participantes = ActividadParticipante.objects.filter(
        actividad=actividad
    ).select_related(
        'habitante',
        'habitante__familia'
    ).order_by('habitante__nombre')
    
    # 2. Inicializamos el libro de openpyxl
    wb = Workbook()
    ws = wb.active
    ws.title = "Participantes en la actividad"
    
    # Habilitar líneas de cuadrícula visibles
    ws.views.sheetView[0].showGridLines = True
    
    # 3. Definición de Estilos Institucionales (Azul y Gris)
    fuente_titulo = Font(name='Arial', size=14, bold=True, color='0F2027')
    fuente_subtitulo = Font(name='Arial', size=10, italic=True, color='555555')
    fuente_cabecera = Font(name='Arial', size=11, bold=True, color='FFFFFF')
    fuente_datos = Font(name='Arial', size=10)
    
    fill_cabecera = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
    fill_cebra = PatternFill(start_color='F2F4F7', end_color='F2F4F7', fill_type='solid')
    
    borde_delgado = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )
    
    # 4. Construcción del Encabezado del Formato
    ws['A1'] = "CONSEJO COMUNAL MANUEL PULIDO MÉNDEZ"
    ws['A1'].font = fuente_titulo
    ws['A2'] = f"ACTIVIDAD: {actividad.nombre_actividad.upper()}"
    ws['A2'].font = Font(name='Arial', size=12, bold=True, color='1F4E78')
    
    # Detalle de las condiciones aplicadas
    resumen_info = f"Categoría: {actividad.get_categoria_enfoque_display()} | Estatus: {actividad.get_estatus_display()}"
    ws['A3'] = resumen_info
    ws['A3'].font = fuente_subtitulo
    ws['A4'] = f"Fecha de exportación: {datetime.now().strftime('%d/%m/%Y')} | Total registros: {participantes.count()}"
    ws['A4'].font = fuente_subtitulo
    
    # Espacio en blanco
    ws.append([]) 
    
    # 5. Cabecera de la Tabla de Datos
    columnas = ['N°', 'Cédula', 'Apellidos y Nombres', 'Teléfono', 'Email', 'Género', 'Familia', 'Parentesco', 'Fecha Registro', 'Observaciones']
    ws.append(columnas)
    
    fila_cabecera = 6
    for col_num, columna in enumerate(columnas, 1):
        celda = ws.cell(row=fila_cabecera, column=col_num)
        celda.font = fuente_cabecera
        celda.fill = fill_cabecera
        celda.alignment = Alignment(horizontal='center', vertical='center')
        celda.border = borde_delgado
    
    # 6. Llenado Lógico de los Registros
    for indice, p in enumerate(participantes, start=1):
        h = p.habitante
        jefe = h.familia.jefe_familia
        
        nombre_completo = f"{h.apellido}, {h.nombre}"
        telefono = getattr(h, 'telefono', '')
        email = getattr(h, 'email', '')
        genero = h.get_genero_display()
        familia_txt = f"{jefe.nombre} {jefe.apellido}" if jefe else "Sin definir"
        parentesco = "Jefe de Familia" if h.es_jefe_familia else "Miembro"
        fecha_registro = p.fecha_registro.strftime('%d/%m/%Y %H:%M')
        observaciones = p.observaciones or ''
        
        fila_datos = [
            indice,
            f"{h.tipo_cedula}-{h.cedula}" if hasattr(h, 'tipo_cedula') and h.tipo_cedula else str(h.cedula),
            nombre_completo,
            telefono,
            email,
            genero,
            familia_txt,
            parentesco,
            fecha_registro,
            observaciones
        ]
        
        ws.append(fila_datos)
        num_fila_actual = ws.max_row
        
        # Aplicamos estilos a las celdas de datos para que se vea limpio
        for col_num in range(1, len(fila_datos) + 1):
            celda = ws.cell(row=num_fila_actual, column=col_num)
            celda.font = fuente_datos
            celda.border = borde_delgado
            
            # Formato cebra intercalado para lectura ágil
            if indice % 2 == 0:
                celda.fill = fill_cebra
                
            # Alineación específica según el tipo de dato
            if col_num in [1, 2, 4, 5, 6, 8, 9]:
                celda.alignment = Alignment(horizontal='center')
            else:
                celda.alignment = Alignment(horizontal='left')

    # 7. Autoajuste automático del ancho de las columnas
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Ignoramos las primeras filas de títulos para que no ensanchen de más la columna A
            if cell.row < 6:
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Ajustes manuales mínimos para columnas largas
    ws.column_dimensions['C'].width = 30  # Nombre
    ws.column_dimensions['I'].width = 18  # Fecha Registro
    ws.column_dimensions['J'].width = 30  # Observaciones

    # 8. Guardado en Buffer RAM y respuesta HTTP directa de descarga
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="Actividad_{actividad.id}_{datetime.now().strftime("%Y%m%d")}.xlsx"'
    
    return response


# ============================================
# Vistas para Dashboard de Actividad
# ============================================

def dashboard_actividad(request):
    total_actividad = Actividad.objects.filter(is_deleted=False).count()
    
    por_estatus = Actividad.objects.filter(is_deleted=False).values('estatus').annotate(
        count=Count('id')
    ).order_by('estatus')
    
    por_categoria = Actividad.objects.filter(is_deleted=False).values('categoria_enfoque').annotate(
        count=Count('id')
    ).order_by('-count')
    
    activos = Actividad.objects.filter(
        is_deleted=False, 
        estatus='activo'
    ).annotate(
        num_participantes=Count('participantes')
    ).order_by('-fecha_creacion')[:10]
    
    total_participaciones = ActividadParticipante.objects.count()
    ultimos_actividad = Actividad.objects.filter(is_deleted=False).order_by('-fecha_creacion')[:5]
    
    hoy = timezone.now().date()
    en_7_dias = hoy + timedelta(days=7)
    proximos_cerrar = Actividad.objects.filter(
        is_deleted=False,
        estatus='activo',
        fecha_fin__range=[hoy, en_7_dias]
    ).order_by('fecha_fin')[:5]
    
    context = {
        'total_actividad': total_actividad,
        'por_estatus': por_estatus,
        'por_categoria': por_categoria,
        'activos': activos,
        'total_participaciones': total_participaciones,
        'ultimos_actividad': ultimos_actividad,
        'proximos_cerrar': proximos_cerrar,
    }
    
    return render(request, 'actividad/dashboard_actividad.html', context)