"""
Vistas para la gestión de censos comunitarios del Consejo Comunal.
"""
from django.shortcuts import render, redirect, get_object_or_404, reverse
from django.contrib import messages
from django.core.paginator import Paginator
from django.db.models import Q, Sum, Count
from django.http import JsonResponse, HttpResponse
from django.utils import timezone
from django.views.decorators.http import require_POST, require_GET
from django.db import transaction
import json
from datetime import datetime, timedelta
import csv
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ✅ CORRECCIÓN: Eliminado 'Person' de las importaciones
from .models import (
    Censo, CensoParticipante, Habitante
)
from .forms import (
    CensoForm, AsignarParticipanteForm
)
from .decorators import admin_required


# ============================================
# Vistas para Censos
# ============================================

def censos(request):
    """
    Vista para listar y buscar censos comunitarios.
    """
    query = request.GET.get('q', '')
    estatus = request.GET.get('estatus', '')
    categoria = request.GET.get('categoria', '')
    
    censos_list = Censo.objects.filter(is_deleted=False).select_related('creado_por')
    
    if query:
        censos_list = censos_list.filter(
            Q(nombre_censo__icontains=query) |
            Q(descripcion__icontains=query)
        )
    
    if estatus:
        censos_list = censos_list.filter(estatus=estatus)
    
    if categoria:
        censos_list = censos_list.filter(categoria_enfoque=categoria)
    
    censos_list = censos_list.annotate(num_participantes=Count('participantes'))
    censos_list = censos_list.order_by('-fecha_creacion')
    
    paginator = Paginator(censos_list, 15)
    page_number = request.GET.get('page')
    censos_page = paginator.get_page(page_number)

    conteos = censos_list.aggregate(
        activo=Count('pk', filter=Q(estatus='activo')),
        cerrado=Count('pk', filter=Q(estatus='cerrado')),
        archivado=Count('pk', filter=Q(estatus='archivado')),
    )
    
    context = {
        'censos': censos_page,
        'query': query,
        'estatus': estatus,
        'categoria': categoria,
        'censo_form': CensoForm(user=request.user),
        'estatus_choices': Censo.ESTATUS_CHOICES,
        'categoria_choices': Censo.CATEGORIA_ENFOQUE_CHOICES,
        'count_activo': conteos['activo'],
        'count_cerrado': conteos['cerrado'],
        'count_archivado': conteos['archivado'],
    }
    
    return render(request, 'censos/censo_list.html', context)


def crear_censo(request):
    if request.method == 'POST':
        form = CensoForm(request.POST, user=request.user)
        if form.is_valid():
            censo = form.save()
            messages.success(request, f'Censo "{censo.nombre_censo}" creado exitosamente.')
            # Si es una petición AJAX, devolver respuesta JSON
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({
                    'success': True,
                    'message': f'Censo "{censo.nombre_censo}" creado exitosamente.',
                    'redirect_url': str(reverse('detalle_censo', kwargs={'pk': censo.pk}))
                })
            return redirect('censos')
        else:
            # Si es una petición AJAX, devolver los errores del formulario en JSON
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                errors = {}
                for field, error_list in form.errors.items():
                    field_label = form.fields[field].label if field in form.fields else field
                    errors[field] = [str(e) for e in error_list]
                return JsonResponse({
                    'success': False,
                    'errors': errors,
                    'message': 'Por favor corrija los errores en el formulario.'
                }, status=400)
            messages.error(request, 'Por favor corrija los errores en el formulario.')
    else:
        form = CensoForm(user=request.user)
    
    context = {'form': form}
    return render(request, 'censos/censo_list.html', context)


def detalle_censo(request, pk):
    """
    Vista para ver el detalle completo de un censo.
    Incluye el listado de participantes registrados.
    """
    # ✅ CORRECCIÓN: Rutas adaptadas al modelo unificado
    censo = get_object_or_404(
        Censo.objects.filter(is_deleted=False).prefetch_related(
            'participantes_censo__habitante',
            'participantes_censo__habitante__familia'
        ).select_related('creado_por'),
        pk=pk
    )
    
    participantes = censo.participantes_censo.all()
    stats_genero = {}
    stats_familias = {}
    
    for p in participantes:
        # ✅ CORRECCIÓN: Acceso directo a campos de Habitante
        genero_display = p.habitante.get_genero_display()
        stats_genero[genero_display] = stats_genero.get(genero_display, 0) + 1
        
        jefe = p.habitante.familia.jefe_familia
        familia_nombre = f"{jefe.nombre} {jefe.apellido}" if jefe else "Sin definir"
        stats_familias[familia_nombre] = stats_familias.get(familia_nombre, 0) + 1
    
    context = {
        'censo': censo,
        'participantes': participantes,
        'asignar_form': AsignarParticipanteForm(censo=censo),
        'stats_genero': stats_genero,
        'stats_familias': stats_familias,
    }
    
    return render(request, 'censos/censo_detail.html', context)


def editar_censo(request, pk):
    censo = get_object_or_404(Censo, pk=pk, is_deleted=False)

    if request.method == "POST":
        form = CensoForm(request.POST, instance=censo, user=request.user)
        if form.is_valid():
            censo = form.save()
            messages.success(request, f'Censo "{censo.nombre_censo}" actualizado exitosamente.')
            # Si es una petición AJAX, devolver respuesta JSON
            if request.headers.get("X-Requested-With") == "XMLHttpRequest":
                return JsonResponse({
                    "success": True,
                    "message": f'Censo "{censo.nombre_censo}" actualizado exitosamente.',
                    "redirect_url": str(reverse("detalle_censo", kwargs={"pk": censo.pk}))
                })
            return redirect("detalle_censo", pk=censo.pk)
        else:
            # Si es una petición AJAX, devolver los errores del formulario en JSON
            if request.headers.get("X-Requested-With") == "XMLHttpRequest":
                errors = {}
                for field, error_list in form.errors.items():
                    field_label = form.fields[field].label if field in form.fields else field
                    errors[field] = [str(e) for e in error_list]
                return JsonResponse({
                    "success": False,
                    "errors": errors,
                    "message": "Por favor corrija los errores en el formulario."
                }, status=400)
            messages.error(request, "Por favor corrija los errores en el formulario.")
    else:
        form = CensoForm(instance=censo, user=request.user)

    context = {"form": form, "censo": censo}
    return render(request, "censos/censo_form.html", context)


@require_POST
def eliminar_censo(request, pk):
    censo = get_object_or_404(Censo, pk=pk, is_deleted=False)
    nombre = censo.nombre_censo
    censo.delete()
    messages.success(request, f'Censo "{nombre}" eliminado exitosamente.')
    return redirect('censos')


@require_GET
def api_censo(request, pk):
    censo = get_object_or_404(Censo, pk=pk, is_deleted=False)
    
    data = {
        'id': censo.id,
        'nombre_censo': censo.nombre_censo,
        'fecha_inicio': censo.fecha_inicio.strftime('%Y-%m-%d'),
        'fecha_fin': censo.fecha_fin.strftime('%Y-%m-%d') if censo.fecha_fin else '',
        'descripcion': censo.descripcion,
        'categoria_enfoque': censo.categoria_enfoque,
        'categoria_enfoque_display': censo.get_categoria_enfoque_display(),
        'estatus': censo.estatus,
        'estatus_display': censo.get_estatus_display(),
        'fecha_creacion': censo.fecha_creacion.strftime('%d/%m/%Y %H:%M'),
        'creado_por': censo.creado_por.username if censo.creado_por else '',
        'participantes_count': censo.cantidad_participantes(),
    }
    
    return JsonResponse(data)


# ============================================
# Vistas para Gestión de Participantes
# ============================================

def buscar_habitantes_censo(request):
    """
    API para buscar habitantes para registrar en un censo.
    """
    query = request.GET.get('q', '')
    censo_id = request.GET.get('censo_id', '')
    
    if len(query) < 2:
        return JsonResponse({'habitantes': []})
    
    # ✅ CORRECCIÓN: Filtros y selects adaptados a Habitante
    habitantes = Habitante.objects.filter(
        is_deleted=False
    ).filter(
        Q(nombre__icontains=query) |
        Q(apellido__icontains=query) |
        Q(cedula__icontains=query)
    ).select_related('familia').order_by('nombre')
    
    if censo_id:
        registrados = CensoParticipante.objects.filter(
            censo_id=censo_id
        ).values_list('habitante_id', flat=True)
        habitantes = habitantes.exclude(id__in=registrados)

    habitantes = habitantes[:20]
    
    resultados = []
    for h in habitantes:
        jefe = h.familia.jefe_familia
        resultados.append({
            'id': h.id,
            'nombre': f"{h.nombre} {h.apellido}",
            'documento': h.cedula,
            'telefono': getattr(h, 'telefono', ''),
            'familia': f"{jefe.nombre} {jefe.apellido}" if jefe else "Sin definir",
            'parentesco': "Jefe de Familia" if h.es_jefe_familia else "Miembro",
            'genero': h.get_genero_display(),
        })
    
    return JsonResponse({'habitantes': resultados})

def asignar_participante(request, pk):
    censo = get_object_or_404(Censo, pk=pk, is_deleted=False)
    
    if request.method == 'POST':
        form = AsignarParticipanteForm(request.POST, censo=censo)
        if form.is_valid():
            participante = form.save()
            # ✅ CORRECCIÓN: Nombres directos del habitante
            nombre_completo = f"{participante.habitante.nombre} {participante.habitante.apellido}"
            messages.success(request, f'{nombre_completo} registrado exitosamente en el censo.')
            return redirect('detalle_censo', pk=censo.pk)
        else:
            messages.error(request, 'Por favor corrija los errores en el formulario.')
    else:
        form = AsignarParticipanteForm(censo=censo)
    
    return render(request, 'censos/censo_detail.html', {'form': form, 'censo': censo})


@require_POST
def remover_participante(request, pk, habitante_id):
    censo = get_object_or_404(Censo, pk=pk, is_deleted=False)
    
    try:
        participante = CensoParticipante.objects.get(
            censo=censo,
            habitante_id=habitante_id
        )
        nombre_completo = f"{participante.habitante.nombre} {participante.habitante.apellido}"
        participante.delete()
        messages.success(request, f'{nombre_completo} removido del censo exitosamente.')
    except CensoParticipante.DoesNotExist:
        messages.error(request, 'El habitante no está registrado en este censo.')
    
    return redirect('detalle_censo', pk=censo.pk)


@require_GET
def api_participante(request, participante_id):
    participante = get_object_or_404(CensoParticipante, id=participante_id)
    
    data = {
        'id': participante.id,
        'censo': participante.censo.id,
        'censo_nombre': participante.censo.nombre_censo,
        'habitante': participante.habitante.id,
        'habitante_nombre': participante.nombre_habitante,
        'habitante_documento': participante.documento_habitante,
        'fecha_registro': participante.fecha_registro.strftime('%d/%m/%Y %H:%M'),
        'observaciones': participante.observaciones or '',
    }
    
    return JsonResponse(data)


# ============================================
# Vistas para Exportación
# ============================================

def exportar_participantes_censo(request, pk):
    """
    Vista para exportar la lista de participantes de un censo a CSV.
    """
    censo = get_object_or_404(Censo, pk=pk, is_deleted=False)
    
    # ✅ CORRECCIÓN: Rutas de relación corregidas
    participantes = CensoParticipante.objects.filter(
        censo=censo
    ).select_related(
        'habitante',
        'habitante__familia'
    ).order_by('habitante__nombre')
    
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="censo_{censo.id}_{datetime.now().strftime("%Y%m%d")}.csv"'
    
    writer = csv.writer(response, delimiter=';', quoting=csv.QUOTE_NONNUMERIC)
    writer.writerow([
        'N°', 'Cédula', 'Nombre Completo', 'Teléfono', 'Email',
        'Género', 'Familia', 'Parentesco', 'Fecha Registro', 'Observaciones'
    ])
    
    for i, p in enumerate(participantes, 1):
        h = p.habitante
        jefe = h.familia.jefe_familia
        
        writer.writerow([
            i,
            h.cedula,
            f"{h.nombre} {h.apellido}",
            getattr(h, 'telefono', ''),
            getattr(h, 'email', ''),
            h.get_genero_display(),
            f"{jefe.nombre} {jefe.apellido}" if jefe else "",
            "Jefe de Familia" if h.es_jefe_familia else "Miembro",
            p.fecha_registro.strftime('%d/%m/%Y %H:%M'),
            p.observaciones or '',
        ])
    
    return response


def exportar_participantes_censo_excel(request, pk):
    """
    Genera un archivo Excel profesional (.xlsx) en memoria RAM con los 
    participantes de un censo comunitario, usando el mismo estilizado 
    que los reportes demográficos.
    """
    censo = get_object_or_404(Censo, pk=pk, is_deleted=False)
    
    participantes = CensoParticipante.objects.filter(
        censo=censo
    ).select_related(
        'habitante',
        'habitante__familia'
    ).order_by('habitante__nombre')
    
    # 2. Inicializamos el libro de openpyxl
    wb = Workbook()
    ws = wb.active
    ws.title = "Participantes Censo"
    
    # Habilitar líneas de cuadrícula visibles
    ws.views.sheetView[0].showGridLines = True
    
    # 3. Definición de Estilos Institucionales (Azul y Gris)
    fuente_titulo = Font(name='Arial', size=14, bold=True, color='0F2027')
    fuente_subtitulo = Font(name='Arial', size=10, italic=True, color='555555')
    fuente_cabecera = Font(name='Arial', size=11, bold=True, color='FFFFFF')
    fuente_datos = Font(name='Arial', size=10)
    
    fill_cabecera = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
    fill_cebra = PatternFill(start_color='F2F4F7', end_color='F2F4F7', fill_type='solid')
    
    borde_delgado = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )
    
    # 4. Construcción del Encabezado del Formato
    ws['A1'] = "CONSEJO COMUNAL MANUEL PULIDO MÉNDEZ"
    ws['A1'].font = fuente_titulo
    ws['A2'] = f"CENSO: {censo.nombre_censo.upper()}"
    ws['A2'].font = Font(name='Arial', size=12, bold=True, color='1F4E78')
    
    # Detalle de las condiciones aplicadas
    resumen_info = f"Categoría: {censo.get_categoria_enfoque_display()} | Estatus: {censo.get_estatus_display()}"
    ws['A3'] = resumen_info
    ws['A3'].font = fuente_subtitulo
    ws['A4'] = f"Fecha de exportación: {datetime.now().strftime('%d/%m/%Y')} | Total registros: {participantes.count()}"
    ws['A4'].font = fuente_subtitulo
    
    # Espacio en blanco
    ws.append([]) 
    
    # 5. Cabecera de la Tabla de Datos
    columnas = ['N°', 'Cédula', 'Apellidos y Nombres', 'Teléfono', 'Email', 'Género', 'Familia', 'Parentesco', 'Fecha Registro', 'Observaciones']
    ws.append(columnas)
    
    fila_cabecera = 6
    for col_num, columna in enumerate(columnas, 1):
        celda = ws.cell(row=fila_cabecera, column=col_num)
        celda.font = fuente_cabecera
        celda.fill = fill_cabecera
        celda.alignment = Alignment(horizontal='center', vertical='center')
        celda.border = borde_delgado
    
    # 6. Llenado Lógico de los Registros
    for indice, p in enumerate(participantes, start=1):
        h = p.habitante
        jefe = h.familia.jefe_familia
        
        nombre_completo = f"{h.apellido}, {h.nombre}"
        telefono = getattr(h, 'telefono', '')
        email = getattr(h, 'email', '')
        genero = h.get_genero_display()
        familia_txt = f"{jefe.nombre} {jefe.apellido}" if jefe else "Sin definir"
        parentesco = "Jefe de Familia" if h.es_jefe_familia else "Miembro"
        fecha_registro = p.fecha_registro.strftime('%d/%m/%Y %H:%M')
        observaciones = p.observaciones or ''
        
        fila_datos = [
            indice,
            f"{h.tipo_cedula}-{h.cedula}" if hasattr(h, 'tipo_cedula') and h.tipo_cedula else str(h.cedula),
            nombre_completo,
            telefono,
            email,
            genero,
            familia_txt,
            parentesco,
            fecha_registro,
            observaciones
        ]
        
        ws.append(fila_datos)
        num_fila_actual = ws.max_row
        
        # Aplicamos estilos a las celdas de datos para que se vea limpio
        for col_num in range(1, len(fila_datos) + 1):
            celda = ws.cell(row=num_fila_actual, column=col_num)
            celda.font = fuente_datos
            celda.border = borde_delgado
            
            # Formato cebra intercalado para lectura ágil
            if indice % 2 == 0:
                celda.fill = fill_cebra
                
            # Alineación específica según el tipo de dato
            if col_num in [1, 2, 4, 5, 6, 8, 9]:
                celda.alignment = Alignment(horizontal='center')
            else:
                celda.alignment = Alignment(horizontal='left')

    # 7. Autoajuste automático del ancho de las columnas
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Ignoramos las primeras filas de títulos para que no ensanchen de más la columna A
            if cell.row < 6:
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Ajustes manuales mínimos para columnas largas
    ws.column_dimensions['C'].width = 30  # Nombre
    ws.column_dimensions['I'].width = 18  # Fecha Registro
    ws.column_dimensions['J'].width = 30  # Observaciones

    # 8. Guardado en Buffer RAM y respuesta HTTP directa de descarga
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="Censo_{censo.id}_{datetime.now().strftime("%Y%m%d")}.xlsx"'
    
    return response


# ============================================
# Vistas para Dashboard de Censos
# ============================================

def dashboard_censos(request):
    total_censos = Censo.objects.filter(is_deleted=False).count()
    
    por_estatus = Censo.objects.filter(is_deleted=False).values('estatus').annotate(
        count=Count('id')
    ).order_by('estatus')
    
    por_categoria = Censo.objects.filter(is_deleted=False).values('categoria_enfoque').annotate(
        count=Count('id')
    ).order_by('-count')
    
    activos = Censo.objects.filter(
        is_deleted=False, 
        estatus='activo'
    ).annotate(
        num_participantes=Count('participantes')
    ).order_by('-fecha_creacion')[:10]
    
    total_participaciones = CensoParticipante.objects.count()
    ultimos_censos = Censo.objects.filter(is_deleted=False).order_by('-fecha_creacion')[:5]
    
    hoy = timezone.now().date()
    en_7_dias = hoy + timedelta(days=7)
    proximos_cerrar = Censo.objects.filter(
        is_deleted=False,
        estatus='activo',
        fecha_fin__range=[hoy, en_7_dias]
    ).order_by('fecha_fin')[:5]
    
    context = {
        'total_censos': total_censos,
        'por_estatus': por_estatus,
        'por_categoria': por_categoria,
        'activos': activos,
        'total_participaciones': total_participaciones,
        'ultimos_censos': ultimos_censos,
        'proximos_cerrar': proximos_cerrar,
    }
    
    return render(request, 'censos/dashboard_censos.html', context)